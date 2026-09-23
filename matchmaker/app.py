"""Site do Matchmaker (Flask): questionário -> tabela ranqueada -> download CSV.

Rodar (da raiz do workspace, para o pacote `finintegra` ser importável):
    python -m matchmaker.app
Depois abra http://localhost:5001
"""
from __future__ import annotations

import csv
import io
import sqlite3
import uuid

from flask import Flask, Response, jsonify, render_template, request

from . import setores
from .recomendador import Criterios, recomendar, _db_path

app = Flask(__name__)

# cache de cidades por UF (a 1ª consulta a um estado guarda a lista)
_CIDADES_UF: dict[str, list[str]] = {}

# cache simples em memória: token -> (criterios_resumo, linhas) para o download.
# Suficiente para uso interno de um usuário; some ao reiniciar o servidor.
_CACHE: dict[str, tuple[dict, list[dict]]] = {}

_COLUNAS = ["empresa", "setor", "cnpj", "cidade", "uf", "porte", "capital",
            "idade", "decisor", "email", "telefone", "site", "score"]


@app.route("/")
def index():
    return render_template("index.html", grupos=setores.grupos(), ufs=setores.UFS)


@app.route("/cidades")
def cidades_view():
    """Lista as cidades (com empresas) dos UFs pedidos: /cidades?ufs=SP,RJ."""
    ufs = [u.strip().upper() for u in request.args.get("ufs", "").split(",") if u.strip()]
    db = _db_path()
    out: dict[str, list[str]] = {}
    if db.is_file():
        con = sqlite3.connect(f"file:{db}?mode=ro", uri=True)
        for uf in ufs:
            if uf not in _CIDADES_UF:
                rows = con.execute(
                    "SELECT DISTINCT cidade FROM empresas WHERE uf=? AND cidade<>'' "
                    "ORDER BY cidade", (uf,)).fetchall()
                _CIDADES_UF[uf] = [r[0].title() for r in rows]
            out[uf] = _CIDADES_UF[uf]
        con.close()
    return jsonify(out)


@app.route("/recomendar", methods=["POST"])
def recomendar_view():
    fm = request.form
    setores_ids = fm.getlist("setores")
    if not setores_ids:
        return render_template("index.html", grupos=setores.grupos(),
                               ufs=setores.UFS,
                               erro="Escolha ao menos um setor para prospectar.")

    ufs = fm.getlist("ufs")
    crit = Criterios(
        setores_ids=setores_ids,
        ufs=ufs,
        cidades_uf=_parse_cidades(fm.getlist("cidade"), ufs),
        portes=fm.getlist("portes"),
        capital_min=_num(fm.get("capital_min"), 0.0),
        capital_max=_num(fm.get("capital_max"), 0.0),
        idade_min=int(_num(fm.get("idade_min"), 0)),
        idade_max=int(_num(fm.get("idade_max"), 0)),
        natureza=fm.get("natureza", ""),
        so_matriz="so_matriz" in fm,
        so_filial="so_filial" in fm,
        so_simples="so_simples" in fm,
        excluir_simples="excluir_simples" in fm,
        so_decisor="so_decisor" in fm,
        exigir_email="exigir_email" in fm,
        exigir_telefone="exigir_telefone" in fm,
        exigir_dominio="exigir_dominio" in fm,
        excluir_mei="excluir_mei" in fm,
        excluir_recjud="excluir_recjud" in fm,
        verificar_site="verificar_site" in fm,
        setores_sec_ids=fm.getlist("setores_sec"),
        n_estab_min=int(_num(fm.get("n_estab_min"), 0)),
        socios_min=int(_num(fm.get("socios_min"), 0)),
        socios_max=int(_num(fm.get("socios_max"), 0)),
        so_estrangeiro="so_estrangeiro" in fm,
        cep_prefixo=fm.get("cep_prefixo", "").strip(),
        divida=fm.get("divida", ""),
        n=int(_num(fm.get("n"), 50)),
    )
    linhas = recomendar(crit)

    reg_parts = []
    for uf in ufs:
        cids = crit.cidades_uf.get(uf.upper(), [])
        reg_parts.append(f"{uf} ({', '.join(cids)})" if cids else uf)
    resumo = {
        "setores": ", ".join(setores.nome_do_id(s) for s in setores_ids),
        "ufs": ", ".join(reg_parts) if reg_parts else "Brasil todo",
        "portes": _portes_label(crit.portes),
        "capital_min": f"R$ {crit.capital_min:,.0f}".replace(",", ".") if crit.capital_min else "sem mínimo",
        "idade_min": f"{crit.idade_min}+ anos" if crit.idade_min else "qualquer",
        "total": len(linhas),
    }
    token = uuid.uuid4().hex[:12]
    _CACHE[token] = (resumo, linhas)
    return render_template("resultados.html", resumo=resumo, linhas=linhas, token=token)


@app.route("/baixar/<token>.csv")
def baixar(token: str):
    item = _CACHE.get(token)
    if not item:
        return "Busca expirada. Refaça a consulta.", 404
    _, linhas = item
    buf = io.StringIO()
    buf.write("﻿")                       # BOM: Excel abre acentos certo
    w = csv.DictWriter(buf, fieldnames=_COLUNAS)
    w.writeheader()
    w.writerows(linhas)
    return Response(
        buf.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f"attachment; filename=leads_{token}.csv"},
    )


_LABELS = {"empresa": "Empresa", "setor": "Setor", "cnpj": "CNPJ", "cidade": "Cidade",
           "uf": "UF", "porte": "Porte", "capital": "Capital", "idade": "Anos",
           "decisor": "Decisor", "email": "E-mail", "telefone": "Telefone",
           "site": "Site", "score": "Score"}
_LARGURAS = {"empresa": 42, "setor": 32, "cnpj": 20, "cidade": 18, "uf": 5,
             "porte": 9, "capital": 12, "idade": 8, "decisor": 28, "email": 32,
             "telefone": 18, "site": 42, "score": 7}


@app.route("/baixar/<token>.xlsx")
def baixar_xlsx(token: str):
    item = _CACHE.get(token)
    if not item:
        return "Busca expirada. Refaça a consulta.", 404
    _, linhas = item
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Empresas"
    ws.append([_LABELS.get(c, c) for c in _COLUNAS])
    for l in linhas:
        ws.append([l.get(c, "") for c in _COLUNAS])
    # cabeçalho em negrito, fundo verde, e congelado
    fill = PatternFill("solid", fgColor="173B32")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions          # filtro do Excel no cabeçalho
    for i, c in enumerate(_COLUNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _LARGURAS.get(c, 15)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return Response(
        buf.read(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=leads_{token}.xlsx"},
    )


def _parse_cidades(valores, ufs):
    """Chips 'UF|Cidade' -> {UF: [cidades]}. Cidade digitada sem UF aplica nos
    UFs marcados (fallback). UF com cidades restringe; UF sem = todas."""
    d: dict[str, list[str]] = {}
    bare: list[str] = []
    for v in valores:
        v = (v or "").strip()
        if not v:
            continue
        if "|" in v:
            uf, cid = v.split("|", 1)
            uf, cid = uf.strip().upper(), cid.strip()
            if uf and cid:
                d.setdefault(uf, []).append(cid)
        else:
            bare.append(v)
    for uf in ufs:
        if bare:
            d.setdefault(uf.strip().upper(), []).extend(bare)
    return d


def _num(v, default):
    try:
        return float(str(v).replace(".", "").replace(",", ".")) if v not in (None, "") else default
    except (ValueError, TypeError):
        return default


def _portes_label(portes: list[str]) -> str:
    mapa = {"01": "ME", "03": "EPP", "05": "Grande"}
    return ", ".join(mapa.get(p, p) for p in portes) if portes else "todos"


def _ip_local() -> str:
    """Descobre o IP da máquina na rede local (pra abrir do celular/outro PC)."""
    import socket
    s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
    try:
        s.connect(("8.8.8.8", 80))          # não envia nada; só descobre a rota
        return s.getsockname()[0]
    except OSError:
        return "127.0.0.1"
    finally:
        s.close()


if __name__ == "__main__":
    ip = _ip_local()
    print("=" * 54)
    print("  Matchmaker no ar. Abra em um navegador:")
    print(f"    - neste PC:            http://localhost:5001")
    print(f"    - celular/outro PC:    http://{ip}:5001")
    print("    (mesma rede Wi-Fi; se não abrir, libere a porta 5001 no firewall)")
    print("=" * 54)
    # host 0.0.0.0 = aceita conexões da rede local, não só deste PC
    app.run(host="0.0.0.0", port=5001, debug=False)
